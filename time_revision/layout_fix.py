from docx import Document
from pathlib import Path
p=Path(__file__).resolve().parent.parent/'redshift_paper/temporal_redshift_paper.docx'
d=Document(p)
for q in list(d.paragraphs):
 if not q.text.strip() and q._p.xpath('.//w:br[@w:type="page"]'):
  q._p.getparent().remove(q._p)
for q in d.paragraphs:
 if q.text in ['Appendix A. Every final-test object','References']:q.paragraph_format.page_break_before=True
d.save(p)
