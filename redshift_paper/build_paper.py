from pathlib import Path
import json, csv, math, hashlib, shutil, zipfile, re
import numpy as np
from scipy.special import lambertw
from scipy.integrate import quad
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

P=Path(__file__).resolve().parent
S=P.parent/'redshift_paper_sources/expanded/time-redshift-expanded'
if not S.exists():
    with zipfile.ZipFile(P.parent/'redshift_paper_sources/time-redshift-expanded.zip') as z:
        z.extractall(S.parent)
fits=json.loads((S/'frozen_parameters.json').read_text())
scores=json.loads((S/'scores.json').read_text())
features=json.loads((S/'features.json').read_text())
predictions=json.loads((S/'frozen_predictions.json').read_text())
test=sorted([r for r in json.loads((S/'scored_predictions.json').read_text()) if r['split']=='test'],key=lambda r:r['distance_mpc'])
K=fits['static_flux']['alpha_per_mpc'];KML=fits['static_flux']['alpha_per_million_ly'];C=299792.458

# Recheck saved results without altering the source experiment.
checks={}
seal=json.loads((S/'prediction_seal.json').read_text())
for fn,key in [('protocol.json','protocol_sha256'),('frozen_parameters.json','parameters_sha256'),('frozen_predictions.json','predictions_sha256')]:
    assert hashlib.sha256((S/fn).read_bytes()).hexdigest()==seal[key]
for m in ['static_flux','flrw']:
    rms=float(np.sqrt(np.mean([(C*(r[m]['predicted_z']-r['observed_z']))**2 for r in test])))
    assert abs(rms-scores['test']['models'][m]['rmse_kms'])<1e-8
    checks[m+'_recomputed_rmse']=rms
for R in [1,10,100,1000]:
    x=KML*R;ue=np.expm1(-x)
    integ=quad(lambda u:1/(1+u),ue,0,epsabs=1e-13)[0]
    assert abs(integ-x)<1e-12
    assert abs(1/(1+ue)-np.exp(x))<1e-12
checks.update(frozen_hashes_match=True,temporal_index_identity=True)
(P/'verification.json').write_text(json.dumps(checks,indent=2))

# Full machine-readable data: every split, no omitted test object.
out=[]
labels={s:json.loads((S/(s+'_labels.json')).read_text()) for s in ['train','validation','test']}
for r in predictions:
    v=labels[r['split']][str(r['pgc'])]
    out.append(dict(pgc=r['pgc'],group_pgc=r['group'],split=r['split'],sky_tile=r['tile'],ra_deg=r['ra'],dec_deg=r['dec'],
                    sbf_modulus_mag=r['dm'],sbf_modulus_error_mag=r['dm_err'],catalog_distance_mpc=r['distance_mpc'],
                    observed_cmb_cz_kms=v,observed_cmb_z=v/C,time_predicted_z=r['static_flux']['predicted_z'],
                    time_conditional_sigma_z=r['static_flux']['sigma_z'],flrw_predicted_z=r['flrw']['predicted_z'],
                    time_residual_kms=C*r['static_flux']['predicted_z']-v))
with (P/'all_164_groups.csv').open('w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(out[0]));w.writeheader();w.writerows(out)

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
d=np.array([r['distance_mpc'] for r in test]);obs=np.array([r['observed_z']*C for r in test]);v=np.array([r['static_flux']['predicted_z']*C for r in test]);vf=np.array([r['flrw']['predicted_z']*C for r in test])
fig,axes=plt.subplots(2,1,figsize=(7,5.5),sharex=True,gridspec_kw={'height_ratios':[2,1]})
axes[0].scatter(d,obs,s=22,color='#202e3b',label='Observed CMB-frame cz')
axes[0].plot(d,v,color='#007c83',label='Temporal transport')
axes[0].plot(d,vf,color='#b57427',ls='--',label='Expansion reference')
axes[0].set_ylabel('Redshift × c (km/s)');axes[0].legend(frameon=False,fontsize=8)
axes[1].axhline(0,color='#aaaaaa',lw=1);axes[1].scatter(d,v-obs,s=19,color='#007c83')
axes[1].set_ylabel('Residual (km/s)');axes[1].set_xlabel('Published SBF distance (Mpc)')
for ax in axes:ax.spines[['top','right']].set_visible(False)
fig.tight_layout();fig.savefig(P/'test_predictions.png',dpi=200);plt.close(fig)

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492)
# standard_business_brief, with a named scientific-paper typography override:
# Cambria body, black 14/12/11pt headings, 22pt title, compact 9pt data tables.
tokens={'preset':'standard_business_brief','override':'scientific_paper','body_font':'Cambria','body_pt':11,
        'body_after_pt':6,'line':1.10,'heading_sizes':[14,12,11],'heading_color':'000000','table_font_pt':9,
        'page_dxa':[12240,15840],'margins_dxa':1440,'table_width_dxa':9360,'table_indent_dxa':120}
(P/'design_tokens.json').write_text(json.dumps(tokens,indent=2))
def style(name,size,bold=False,before=0,after=6):
    s=doc.styles[name];s.font.name='Cambria';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        s.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)
    for border in list(s.element.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
    s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after);s.paragraph_format.line_spacing=1.10
    s.paragraph_format.widow_control=True
    return s
style('Normal',11);style('Title',22,True,0,8);style('Subtitle',11,False,0,8)
style('Heading 1',14,True,16,8);style('Heading 2',12,True,12,6);style('Heading 3',11,True,8,4)
style('Caption',9,False,4,4)
for name,size in [('Equation',11),('Data',9),('Reference',9)]:
    doc.styles.add_style(name,1);style(name,size,False,4 if name=='Equation' else 0,4)
for name in ['Heading 1','Heading 2','Heading 3']:doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Equation'].font.name='Cambria Math'
hp=sec.header.paragraphs[0];hp.text='TEMPORAL TRANSPORT  |  RESEARCH HYPOTHESIS';hp.style=doc.styles['Caption']
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=fp.add_run('Page ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
def put_text(para,t):
    for part in re.split('(Dᴸ|Dᴬ)',t):
        if part in ['Dᴸ','Dᴬ']:
            para.add_run('D');para.add_run('L' if part=='Dᴸ' else 'A').font.subscript=True
        else:para.add_run(part)
def p(t,sty=None):
    para=doc.add_paragraph(style=sty);put_text(para,t);return para
def h(t):return doc.add_heading(t,level=1)
def eq(t):
    para=p('', 'Equation');para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    from docx.enum.text import WD_TAB_ALIGNMENT
    para.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
    match=re.search(r'\((\d+)\)\s*$',t);label=match.group(0);t=t[:match.start()].rstrip()
    t=re.sub(r' {3,}','   ',t)
    om=OxmlElement('m:oMath')
    def mr(txt):
        rr=OxmlElement('m:r');tt=OxmlElement('m:t');tt.text=txt;rr.append(tt);return rr
    for part in re.split('(Dᴸ|Dᴬ)',t):
        if part in ['Dᴸ','Dᴬ']:
            sub=OxmlElement('m:sSub');base=OxmlElement('m:e');base.append(mr('D'));idx=OxmlElement('m:sub');idx.append(mr('L' if part=='Dᴸ' else 'A'));sub.append(base);sub.append(idx);om.append(sub)
        else:om.append(mr(part))
    para._p.append(om);para.add_run('\t'+label)
    return para
def table(headers,rows,widths):
    assert sum(widths)==9360
    tab=doc.add_table(rows=1,cols=len(headers));tab.autofit=False
    pr=tab._tbl.tblPr
    for tag in ['tblW','tblInd','tblLayout','tblCellMar','tblBorders']:
        for el in list(pr.findall(qn('w:'+tag))):pr.remove(el)
    w=OxmlElement('w:tblW');w.set(qn('w:w'),'9360');w.set(qn('w:type'),'dxa');pr.append(w)
    ind=OxmlElement('w:tblInd');ind.set(qn('w:w'),'120');ind.set(qn('w:type'),'dxa');pr.append(ind)
    layout=OxmlElement('w:tblLayout');layout.set(qn('w:type'),'fixed');pr.append(layout)
    mar=OxmlElement('w:tblCellMar')
    for tag,val in [('top',80),('bottom',80),('start',120),('end',120)]:
        el=OxmlElement('w:'+tag);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');mar.append(el)
    pr.append(mar);borders=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+tag);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D4D4D4');borders.append(el)
    pr.append(borders)
    grid=tab._tbl.tblGrid
    for el in list(grid):grid.remove(el)
    for wd in widths:
        el=OxmlElement('w:gridCol');el.set(qn('w:w'),str(wd));grid.append(el)
    for values in rows:tab.add_row()
    for i,values in enumerate([headers]+rows):
        row=tab.rows[i]
        nr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(nr)
        if i==0:
            rep=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(rep)
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]/1440);tcw=cell._tc.get_or_add_tcPr().find(qn('w:tcW'));tcw.set(qn('w:w'),str(widths[j]));tcw.set(qn('w:type'),'dxa')
            cell.text='';put_text(cell.paragraphs[0],str(value))
            for pp in cell.paragraphs:
                pp.style=doc.styles['Data'];pp.paragraph_format.space_after=Pt(1)
                for rr in pp.runs:rr.bold=i==0
            if i==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'F2F4F7');cell._tc.get_or_add_tcPr().append(fill)
    p('', 'Caption')
    return tab

p('Temporal Transport in a Non-Expanding Universe','Title')
p('An exponential redshift law, a conditional physical mechanism, and a reproducible galaxy-group test','Subtitle')
p('Research hypothesis • 8 September 2026','Caption')
h('Abstract')
p('We investigate a universe with fixed spatial geometry in which a propagation process, rather than increasing source separation, produces a systematic spectral redshift. A local, multiplicative temporal-transfer law gives 1+z=exp(κr), where r is geometric path length. We derive the exponential from continuity and composition and construct a conditional realization using a homogeneous, nondispersive optical response that evolves linearly with time. This realization stretches both wave periods and pulse intervals, but its constitutive physics, matter coupling and gravitational backreaction remain unspecified. A preserved Cosmicflows-4 surface-brightness-fluctuation experiment contains 164 galaxy-group representatives, excluding all groups from an earlier pilot: 104 training, 35 validation and 25 final-test objects. The fitted coefficient is κ=7.7314966×10⁻⁵ per million light-years. With a stated brightness correction, final-test redshift RMSE is 412.441 km/s in cΔz units, compared with 411.934 km/s for a specified expanding reference. The difference is unresolved. We provide every final-test prediction, all 164 records in machine-readable form, the original source catalog and executable analysis. The observations establish limited predictive compatibility, not a unique cause of redshift. Galaxy dynamics are reserved for a separate paper.')
p('Scope. The non-expanding setting is the working hypothesis of this paper. Observational numbers are real catalog values and reproducible calculations; no fictional measurements are introduced as evidence. “Proof” requires discriminating tests beyond the current nearby sample.','Caption')
h('1. Observable and physical distinction')
p('Define the spectroscopic redshift through the frequency of a calibrated emitted spectral feature and its measured arrival frequency, 1+z=νₑ/νₒ. The proposed temporal-transfer factor S is this ratio. A stronger hypothesis states that the duration of a sufficiently short, otherwise unchanged signal is multiplied by the same S. Spectral redshift, transient-duration stretching and a local clock-rate difference are distinct observables. This distinction is essential: an equation for arriving wave periods is not itself a measurement of clocks in intergalactic space.')
p('The empirical experiment uses surface-brightness-fluctuation (SBF) distance estimates for galaxies, not distances and redshifts of individually resolved stars. SBF estimates distance through fluctuations in unresolved stellar light; its calibration and corrections must be retained explicitly [1,2]. The present calculation treats the published SBF distance as a luminosity-distance input. It does not reconstruct the images or the calibration under the proposed propagation law. The catalog redshift is expressed in the CMB frame; cΔz is used solely as a convenient unit for prediction error.')
h('2. Why an exponential law follows')
p('Assume that transfer through adjacent homogeneous path segments composes multiplicatively, that a zero-length segment has no effect, and that the response is continuous. If the same local rule applies everywhere over the sampled range, the transfer factor obeys S(r₁+r₂)=S(r₁)S(r₂), with S(0)=1 and S(r)>0. Taking the logarithm converts this into a continuous additive relation. Its solution is ln S=κr, where κ is a constant inverse length:')
eq('d ln S / dr = κ,          S(r) = 1 + z = exp(κr).                                      (1)')
p('Equivalently, each small segment changes an interval by a fraction κ dr of its current value, rather than by the same absolute number of seconds. Multiplication over many segments gives the exponential. These assumptions explain the functional form conditionally; they do not predict κ, establish the composition rule experimentally, or identify the underlying field. An environmental extension replaces κr by an integral of κ(x,t,direction) along the ray and requires additional measured inputs.')
eq('κ = 7.7314965955 × 10⁻⁵ Mly⁻¹ = 2.5216769238 × 10⁻⁴ Mpc⁻¹.                       (2)')
p('Here one Mly is one million light-years, and r is a length defined using present-day units. The model’s photon travel time need not equal r divided by the present photon speed. The equivalent fitted slope is cκ=75.5979723 km s⁻¹ Mpc⁻¹; this is a dimensional conversion, not an expansion rate in the proposed universe. The displayed digits preserve the computation and should not be read as comparable physical precision.')
table(['Geometric path (Mly)','S = 1+z','One emitted second arrives as'],[[str(R),f'{math.exp(KML*R):.9f}',f'{math.exp(KML*R):.9f} s'] for R in [1,10,100]],[2500,2860,4000])
p('Table 1. Model predictions, not observations of local clocks. Equal stretching of a whole signal is an additional requirement, satisfied by the conditional construction below.','Caption')
h('3. A conditional mechanism: evolving optical response')
p('A stationary local clock field cannot generate the proposed accumulated shift simply by adding contributions from successive wells. For clocks at rest in a stationary metric, the gravitational frequency ratio depends on the endpoint lapse factors; intermediate factors cancel [5]. We therefore consider a genuinely time-dependent propagation response. Laboratory time-refraction research establishes the relevant optical principle: in a spatially homogeneous but temporally changing medium, wavevector is conserved while frequency can change [4]. That principle does not establish an intergalactic medium of the required kind.')
p('As a constructive model, let material rulers define fixed Euclidean positions and material clocks define time t. Introduce a positive, spatially homogeneous effective optical index n(t), assumed nondispersive over the frequencies tested. Let c₀ be the present photon speed, normalize n(tₒ)=1 at the observation epoch, and assume that the calibrated source transition frequency does not change with n. Spatial homogeneity conserves the wavevector q in the adiabatic geometric-optics limit, while the instantaneous dispersion relation and ray speed are:')
eq('ω(t) = c₀ |q| / n(t),                  vγ(t) = c₀ / n(t).                              (3)')
p('Consequently S=n(tₒ)/n(tₑ). The spatial separation r of stationary emitter and observer equals the integral of c₀/n(t) from emission to observation. Choose the index to change linearly, with a positive constant slope γ:')
eq('dn/dt = γ,                  n(t) = 1 + γ(t − tₒ),                  γ = c₀κ.            (4)')
p('The path integral is then exactly r=(c₀/γ) ln[n(tₒ)/n(tₑ)]. Combining it with the frequency ratio gives Equation (1). The same model supplies envelope stretching: differentiate the fixed-distance ray integral for two neighboring emitted pulses. Its endpoint terms give dtₒ/n(tₒ)=dtₑ/n(tₑ), hence dtₒ/dtₑ=n(tₒ)/n(tₑ)=S. This is a local pulse-duration relation, valid when the optical response varies negligibly across each pulse. The derivation links frequency and temporal stretching without requiring spatial expansion.')
eq('r = (c₀/γ) ln(nₒ/nₑ),                 νₑ/νₒ = dtₒ/dtₑ = exp(κr).                 (5)')
p('In this construction the spatial wavevector remains fixed during homogeneous propagation. It is the temporal period that grows. A wavelength redshift inferred by comparison with a present-day laboratory standard remains consistent with the frequency ratio; one must not simultaneously assume that the wave’s spatial wavelength grew as it would in an expanding geometry. Material standards and their coupling to the optical field are explicit assumptions.')
p('A possible effective-field motivation is an approximately force-free, homogeneous scalar n with positive kinetic term Kₙ[(∂ₜn)²−vₙ²|∇n|²]/2 and an almost flat potential V(n). In the absence of significant backreaction its equation gives ∂ₜ²n≈0, admitting Equation (4). Couple its optical response through ε=ε₀n and μ=μ₀n; then the photon speed is c₀/n and impedance is constant. A corresponding electromagnetic Lagrangian density is ε₀nE²/2−B²/(2μ₀n). This is a schematic constitutive model, not a completed cosmological action. Electromagnetic coupling sources the scalar, so constant γ is an approximation that requires a controlled backreaction calculation.')
p('The coefficient γ is an integration constant in this construction. Its fitted value, 2.44996×10⁻¹⁸ s⁻¹ today, is not derived from known fundamental constants. Photon energy decreases as E∝1/n and must be transferred to the response field or its driving sector. Photon number conservation and negligible scattering are assumed separately. Including that sector is necessary for energy accounting, but its stress-energy and the possibility of maintaining static geometry have not been solved. The construction must not silently reintroduce an untested gravitating component as an explanation of galaxy dynamics.')
h('4. Brightness and prediction without using the target redshift')
p('For isotropic emission into Euclidean ray geometry, assume no absorption, no photon creation and no scattering out of the beam. Each arriving photon has energy smaller by S, and equal envelope stretching reduces the photon arrival rate by another factor S. If L is the emitted bolometric luminosity per source-clock second, the received bolometric flux is F=L/(4πr²S²). Defining luminosity distance by F=L/(4πDᴸ²) gives Dᴸ=rS. This derives the brightness relation used in the archived experiment under the stated optical and source assumptions, not from an expansion law.')
eq('Dᴸ = r exp(κr),        κr = W(κDᴸ),        zpred = exp[W(κDᴸ)] − 1.                 (6)')
p('W is the principal real Lambert function, defined by W(x) exp[W(x)]=x. For positive distance and κ, Equation (6) is single-valued. The prediction uses only the measured distance and fitted coefficient: no observed target redshift appears on its right-hand side. Treating the brightness distance as the geometric path instead would give exp(κDᴸ)−1, a different approximation retained only as a diagnostic.')
p('If photon energies redshift but whole signals do not stretch, luminosity distance would instead be r√S in this same geometry. Opacity, dispersion, source-clock evolution or different angular transport would modify the relation again. Thus the successful fit is conditional on a specific brightness prescription. Published SBF bandpass and population corrections were not rebuilt from the images; their use limits the interpretation even though the model inversion itself does not feed the target redshift back into the prediction.')
h('5. Data, selection and fitting')
p('The source is Cosmicflows-4 table 2, distributed through CDS [1]. Select the method-specific SBF modulus μ and its uncertainty σμ, convert Dᴸ=10^[(μ−25)/5] Mpc, require 10≤Dᴸ≤150 Mpc and 0<σμ≤0.30 mag, and exclude all 68 dominant-PGC groups present in the earlier pilot. These cuts produce 213 galaxy measurements. Retaining the smallest-uncertainty SBF measurement in each group, with PGC as the tie-breaker, leaves 164 group representatives. Selection uses no cut on target redshift, but group construction and distance calibration have inherited catalog assumptions.')
p('Divide the sky into 24 right-ascension bins and four equal sin(declination) bins. A fixed SHA256 mapping of each occupied sky tile assigns the split without retries. This gives 104 training groups in 36 tiles, 35 validation groups in 16 tiles and 25 final-test groups in 13 tiles. All parameter fits use training labels only. Validation is diagnostic, not a tuning or selection stage. Parameters, predictions and checksums were saved before held-out scoring in the archived run. This is a computational holdout of published data, not independently blinded observations; neighboring structures and shared calibrations can correlate different tiles.')
p('Write yᵢ=czᵢ in km/s and μ̂ᵢ=c zpred(Dᵢ). Each fitted model minimizes the sum of Gaussian negative log likelihoods, including the logarithm of its distance-dependent variance:')
eq('NLL = Σᵢ { ln σᵢ + (yᵢ − μ̂ᵢ)²/(2σᵢ²) + ½ ln(2π) }.                              (7)')
eq('σᵢ² = (300 km/s)² + [(∂μ̂ᵢ/∂Dᵢ) σD,ᵢ]²,      σD,ᵢ = Dᵢ ln(10) σμ,ᵢ/5.           (8)')
p('The fixed 300 km/s floor represents unresolved motions phenomenologically. The individual representative’s systemic CMB-frame velocity is used; internal group motion is not removed. The method omits correlated flows, common distance-scale errors, selection-function corrections and a latent-distance likelihood. This is an explicit approximate error model. For the temporal prescription, ∂μ̂/∂D=cκ/[1+W(κD)], which is used in the variance rather than a derivative evaluated with the observed target redshift.')
p('The comparator is a specified flat FLRW distance law with matter fraction 0.3, cosmological-constant fraction 0.7 and negligible radiation [3]. It fits only H₀, while the temporal model fits only cκ. Its luminosity distance is (c/H₀)(1+z) times the integral from 0 to z of [0.3(1+u)³+0.7]⁻¹ᐟ²; numerical inversion predicts z from the same input distance. The fitted comparator scale is H₀=75.8848424 km s⁻¹ Mpc⁻¹. This is a controlled benchmark, not an analysis of all cosmological observations.')
h('6. Results and what they establish')
table(['Prescription','Validation RMSE','Test RMSE','Test MAE'],[
 ['Temporal transport + brightness correction','435.927','412.441','336.927'],
 ['Specified expanding reference','435.900','411.934','336.252'],
 ['Distance-as-path exponential approximation','436.301','414.627','339.437'],
 ['Temporal model, earlier pilot rate fixed','444.872','427.645','355.102']],[4320,1680,1680,1680])
p('Table 2. Errors are in km/s as cΔz, not inferred physical recession speeds. Every prescription is scored on the same held-out objects.','Caption')
p('The expanding reference has a 0.507 km/s lower final-test RMSE. A paired 5,000-draw sky-tile bootstrap for temporal-minus-reference RMSE gives a 95% interval from −0.814 to +1.948 km/s. It includes zero; the sample does not resolve a preferred model. The two laws share the leading low-distance relation z≈κDᴸ, making their small differences difficult to distinguish from distance errors and motions. The temporal test error is approximately 0.001376 in redshift units.')
p('Training-only sky-tile bootstrap refits give a conditional 95% range cκ=73.053–78.656 km s⁻¹ Mpc⁻¹. It excludes neither common calibration errors nor the consequences of model misspecification. The earlier pilot rate was 69.5192 km s⁻¹ Mpc⁻¹; the new rate is 8.74% higher. A universal coefficient is therefore not established. Nominal 95% prediction bands cover 22 of 25 test objects, or 88%; they omit parameter uncertainty and cannot be described as validated 95% coverage.')
p('For this paper, the archived parameter and prediction hashes were checked and both primary test RMSE values were independently recomputed from the saved full-precision predictions and labels. No new fit was substituted and no final-test object was removed. Figure 1 and Appendix A show all 25 test objects. The preserved pipeline and the 164-row supplementary table allow independent recalculation; they do not turn exposed observations into a fresh holdout.')
doc.add_picture(str(P/'test_predictions.png'),width=Inches(6.3))
p('Figure 1. Every final-test observation and the two primary predictions (upper panel), with temporal-model prediction minus observation below. The prediction curves nearly overlap. A good distance trend alone does not identify its physical cause.','Caption')
h('7. Distinguishing predictions and required evidence')
p('A convincing test must address consequences that were not used to fit κ. The current 25 final-test objects are exposed and must be treated as development data in any future model change. A new sample should be specified before labels are inspected, with independent distance anchors, forward-modeled photometry and common calibration covariance. Longer-distance predictions below are extrapolations of the hypothesis, not measurements or validated extensions of this fit.')
p('Temporal and spectral agreement. With source evolution modeled independently, a matched transient or periodic signal should have duration stretch equal to 1+z. Multi-band line ratios should give the same z because the proposed response is nondispersive. Chromatic residuals, pulse broadening beyond the source model, polarization effects or angular blurring would require additional response physics or reject the simplest construction. Measuring transient dilation alone would not distinguish this model from expansion; it instead tests its own brightness assumption.')
p('Distance duality and surface brightness. In the proposed Euclidean ray geometry, angular-diameter distance Dᴬ=r and Dᴸ=(1+z)Dᴬ. Bolometric surface brightness of otherwise identical resolved sources scales as (1+z)⁻². Standard metric, photon-conserving distance duality has Dᴸ=(1+z)²Dᴬ, giving a different comparison [3]. The temporal optical construction changes the propagation assumptions, so its relation must be tested rather than substituted into standard distance reductions. Independent rulers, fluxes and careful source-evolution corrections are required; SBF distance alone cannot perform this test.')
p('Epoch drift and local propagation. With constant γ and stationary source separation r, the derived redshift is independent of observation epoch: dz/dtₒ=0. This is a specific prediction of the homogeneous construction, not of every temporal theory. The local photon speed instead changes at d ln vγ/dt=−γ/n, or about −7.7315×10⁻¹¹ per year today relative to fixed material standards. Meaningful laboratory tests compare independent clocks and rulers; they cannot simply interpret a change in a unit-defined value of c. If screening is invoked to suppress local changes, its equations must be specified and the distance law rederived through screened regions.')
p('Global consistency. With n(tₒ)=1, linear evolution reaches n=0 approximately 12.934 billion material-clock years into the past. Before that boundary the construction is undefined; past n<1 also implies a photon speed greater than today’s c₀. This is a limitation of the optical ansatz, not a measured age of the universe. Stellar ages, causal structure, source spectra, background-radiation spectrum and thermal history must be checked in the same matter-clock convention. A theory with the same stars and laboratory physics must explain why the required optical response does not change stellar emission and atomic standards. Neither the assumed background geometry nor these consistency conditions has been established.')
table(['Assumed observed z','Predicted r (Mly)','Predicted Dᴸ (Mly)','Duration factor'],[[str(z),f'{math.log1p(z)/KML:,.1f}',f'{(1+z)*math.log1p(z)/KML:,.1f}',f'{1+z:.2f}'] for z in [.01,.1,.5,1]],[1900,2500,2760,2200])
p('Table 3. Conditional forecast grid from the fitted coefficient. Values beyond the nearby calibration regime are untested. r is geometric distance, not c₀ times lookback time.','Caption')
p('A falsifiable evaluation should preregister the competing transport models, distance range, calibration treatment, sky/group exclusions and prediction metric. Release per-object distances, errors, redshifts, fluxes, bandpasses, selection flags and covariances; preserve frozen parameters and target-free predictions before scoring. Improvement must transfer to that new sample without per-object adjustments. Agreement should include calibrated uncertainty coverage and the independent temporal, angular and spectral tests above, rather than only a reduced average redshift residual.')
h('8. Relationship to galaxy rotation')
p('Galaxy rotation is not included as supporting evidence for this redshift mechanism. Exploratory radial acceleration prescriptions can be written using the dimensional scale c²κ, but a free coupling absorbs changes in κ and prevents an independent numerical identification. A radial fit also does not by itself define a single-valued three-dimensional clock field. The evolving optical response derived here concerns photon propagation and has not supplied the gravitational field equation governing stellar orbits. A combined treatment must derive both sectors, conserve energy consistently and predict vertical motion and lensing before the rotation data can be used as evidence of unification.')
h('9. Conclusion')
p('The exponential law follows from specified composition assumptions and can be realized conditionally by a linearly evolving optical response in fixed spatial geometry. That construction explains how both frequency and pulse intervals could transform, while leaving the numerical rate and the required constitutive physics to be established. The preserved nearby-galaxy experiment shows predictive compatibility with a specified expanding reference, not superiority or proof of non-expansion. Its chief scientific product is a reproducible, falsifiable hypothesis with explicit brightness assumptions, complete test data and distinct observational consequences. The next advance must resolve those consequences and the model’s matter and field consistency, rather than attach an untested explanation of galaxy rotation.')

doc.add_page_break();h('Appendix A. Every final-test object')
p('Distances are the published SBF luminosity-distance inputs. Residual is c(zpred−zobs). The table is rounded for reading; the supplement retains full precision, group IDs, sky positions, split assignments and prediction uncertainties. No final-test object is omitted.','Caption')
table(['PGC','Dᴸ (Mpc)','σμ (mag)','Predicted z','Observed z','Residual (km/s)'],[[str(r['pgc']),f'{r["distance_mpc"]:.3f}',f'{r["dm_err"]:.3f}',f'{r["static_flux"]["predicted_z"]:.7f}',f'{r["observed_z"]:.7f}',f'{C*(r["static_flux"]["predicted_z"]-r["observed_z"]):+.1f}'] for r in test],[1000,1400,1200,1900,1900,1960])
h('Appendix B. Data release and reproduction')
p('The companion redshift_paper_data.zip contains all_164_groups.csv, the original time-redshift-expanded.zip, the paper-generation script, verification results and the derived figure. The original archive includes the full source catalog and field documentation, earlier-pilot group exclusions, protocol, separate label files, frozen parameters and predictions, checksums, scoring code and all results. The CSV combines these records for scientific inspection; that combined file is not a blinded input for a new test.')
p('In an extracted copy of the original archive, use Python with NumPy and SciPy and execute run_test.py with stages prepare, train, predictions, verify and score in that order, followed by make_report.py. These commands regenerate the original experiment outputs in that copy; preserve the distributed archive unchanged. The model uses distances in Mpc internally, κ in Mpc⁻¹, and c=299792.458 km/s. Reproduction recovers a computational split, not an untouched scientific test.')
p('For a single distance D, compute w=W(κD), then zpred=exp(w)−1 and r=D/(1+zpred). For the optical construction, compute nₑ=exp(−κr) at nₒ=1 and elapsed flight time Δt=[1−exp(−κr)]/(c₀κ). Verify the path integral and interval Jacobian against Equation (5). These identities were numerically checked for geometric paths of 1, 10, 100 and 1,000 Mly; they validate algebra and implementation, not the existence of the proposed response field.')
p('A full evidential release for the next experiment must add raw or reprocessable SBF photometry, zero-point and population calibration, bandpass/extinction corrections, covariance and selection functions, and independently measured transient/ruler/environment data. Those additional measurements are not present in this release. In particular, no void-density or potential information was used to fit the reported redshift coefficient.')
h('References')
refs=[
'[1] Tully, R. B., et al. Cosmicflows-4. Astrophysical Journal 944, 94 (2023). arXiv:2209.11238. Catalog: CDS J/ApJ/944/94, table2.dat. https://arxiv.org/abs/2209.11238',
'[2] Blakeslee, J. P., Jensen, J. B., Ma, C.-P., Milne, P. A., and Greene, J. E. The Hubble Constant from Infrared Surface Brightness Fluctuation Distances. Astrophysical Journal 911, 65 (2021). https://arxiv.org/abs/2101.02221',
'[3] Hogg, D. W. Distance measures in cosmology (1999; revised 2000). https://arxiv.org/abs/astro-ph/9905116',
'[4] Broadband frequency translation through time refraction in an epsilon-near-zero material. Nature Communications 11 (2020). DOI:10.1038/s41467-020-15682-2. https://www.nature.com/articles/s41467-020-15682-2',
'[5] Tong, D. General Relativity, chapter 1: Geodesics in Spacetime, gravitational time dilation and redshift. https://www.davidtong.org/teaching/general-relativity/grhtml/S1',
'[6] Preserved analysis: Time replacing expansion: expanded holdout test, 8 September 2026. Distributed source archive time-redshift-expanded.zip; not an externally peer-reviewed publication.'
]
for ref in refs:p(ref,'Reference')
doc.core_properties.title='Temporal Transport in a Non-Expanding Universe'
doc.core_properties.subject='Conditional redshift mechanism and reproducible empirical test'
doc.core_properties.author=''
doc.save(P/'temporal_redshift_paper.docx')
print('Created',P/'temporal_redshift_paper.docx')
