from pathlib import Path
from shutil import copyfile
from docx import Document

ROOT=Path(__file__).resolve().parent.parent
path=ROOT/'redshift_paper/temporal_redshift_paper.docx'
doc=Document(path)
assert not any(p.text.startswith('Appendix D.') for p in doc.paragraphs)
copyfile(path,ROOT/'paper_fixed_atom_update/before.docx')
original=[p._p.xml for p in doc.paragraphs]
def p(t,style='Normal'):doc.add_paragraph(t,style)
def h(t):p(t,'Heading 2')
def e(t):p(t,'Equation')

doc.add_page_break()
p('Appendix D. Fixed atomic standards and an unscreened propagation field','Heading 1')
p('Status: theoretical extension with synthetic checks, 8 September 2026. This branch retains fixed atomic transition frequencies, masses and ruler lengths as explicit assumptions. It introduces no density screening. It does not yet derive atomic invariance from a microscopic interaction, and no new observational fit is claimed. Earlier screened and varying-atom branches remain separate candidates.')
h('D.1 Postulates and derived observables')
p('In a static matter background, let atomic clocks measure t and let a positive field n(x,t), normalized to one today, control traveling photons. Postulate the photon Hamiltonian Hγ=c₀|p|/n. Homogeneity conserves momentum, giving frequency ω=c₀|k|/n and propagation speed cγ=c₀/n. Temporal frequency conversion in engineered media is a physical analogue [18,19], not evidence that intergalactic vacuum follows this law. In an inhomogeneous field, d ln ω/dt=−∂t ln n along a ray; static spatial variation alone does not produce cumulative frequency change between fixed clocks in this model.')
p('Give n a positive kinetic and gradient energy with coefficient K>0, disturbance speed vₙ and potential V. The field Lagrangian density is:')
e('ℒₙ = (K/2)[(∂t n)² − vₙ²|∇n|²] − V(n).')
p('For homogeneous n, constant V and negligible radiation backreaction, the equation is n̈=0. Its solution is n(t)=1+γt, where γ is an initial velocity rather than a predicted constant. For reception today and a source at fixed geometric separation R:')
e('R = ∫[tₑ,0] c₀ dt/(1+γt) = (c₀/γ) ln[1/n(tₑ)].')
e('1+z = nᵣ/nₑ = exp(γR/c₀);     dtᵣ/dtₑ = 1+z.')
p('The second equality follows by differentiating the same ray integral at fixed R. Thus unchanged atomic clocks can register both spectral redshift and longer arrival intervals. With γ/c₀=0.000077315 per million light-years, this recovers the supplied exponential law; today γ=7.7315×10⁻¹¹ yr⁻¹. The coefficient and direction of field motion remain inputs. Numerical integration reproduces the free-field exponential to relative error below 7×10⁻¹⁶. Two independently propagated pulses at γR/c₀=0.1 reproduce their predicted stretch 1.1051709181 to relative error 1.4×10⁻¹¹. These are mathematical checks, not detections.')

doc.add_page_break()
h('D.2 Field choice, energy exchange and extrapolation')
p('The dynamical choice differs from a canonical free field χ=ln n. Canonical χ gives n=exp(γt) and 1+z=1+γR/c₀; canonical n gives the exponential distance law above. Written in χ, the latter has positive kinetic coefficient K exp(2χ), so an unstable potential is not universally required to produce affine n. A field redefinition must transform the kinetic term. Neither choice has yet been uniquely selected by first principles or observations.')
p('The free affine history reaches n=0 at t=−1/γ, about 12.9 billion coordinate years before the reference epoch. The photon Hamiltonian is singular there. This marks the boundary of that approximation, not a derived Big Bang or universe age. Energy exchange already changes the extrapolation. In a homogeneous fixed volume, adiabatically conserved radiation occupations give energy density uγ=A/n. With no additional matter exchange:')
e('K n̈ + V′(n) = uγ/n;     ℰ = (K/2)ṅ² + V(n) + A/n = constant.')
p('Radiation losing energy supplies the field with an equal gain. For V=0 this accelerates n, so the exact exponential requires negligible backreaction. Let η=uγ,0/(Kγ²). At dimensionless distance γR/c₀=7, the free stretch is 1096.633. Conservative simulations give 1095.440 for η=10⁻⁶, 986.314 for η=10⁻⁴, and 459.162 for η=10⁻³. Total energy is conserved to relative error below 1.6×10⁻¹⁵ across the tested runs. Their energy scales are synthetic and not calibrated to gravity or measured radiation densities.')
p('Positive kinetic energy gives the analytic maximum stretch (1+z)max=1+1/(2η) in this reduced bath model. A past turning point of n replaces the free-history singularity and can eventually reverse the sign of the observed shift. Obtaining a stretch of 1100.7 requires η≤0.0004547, or present field kinetic energy at least 1099.7 times the modeled bath energy. This necessary condition does not ensure an accurate exponential. The gravitational consequences of that reservoir remain unspecified.')
h('D.3 Thermal radiation')
p('An already thermal bath remains Planckian under adiabatic, frequency-independent mode rescaling with conserved occupation, giving Tᵣ=Tₑ/(1+z) in the stipulated atomic standards. Since cγ=c₀/n, the fixed-volume mode density gives uγ∝n³T⁴∝1/n, consistently with the energy budget. Stretching does not create a Planck distribution from arbitrary radiation. The FIRAS comparison in C.4 cannot independently determine initial temperature and stretch: different pairs produce the same final spectrum. No new CMB temperature, thermalization mechanism, angular structure or preferred-axis prediction follows here.')

doc.add_page_break()
h('D.4 What must be tested before stronger claims')
p('Atomic invariance is the central unfinished derivation. Traveling photons and electromagnetic binding belong to one underlying theory. Separately stipulating fixed atomic levels and a modified free-photon Hamiltonian does not prove consistency of emission, absorption, magnetic interactions and quantum corrections. A microscopic completion must demonstrate exact protection through a symmetry, or calculate residual atomic effects for comparison with measurements. Related scalar couplings illustrate why this issue is nontrivial [20]. The effective model also selects a preferred frame and lacks a completed gravitational sector.')
p('The first observational priority is a comparison of an optical cavity with an atomic clock. With fixed mirror spacing and the same nondispersive propagation law, cavity frequency is proportional to cγ while the atomic reference is constant. The model predicts d ln(νcavity/νatom)/dt=−γ today, approximately −7.73×10⁻¹¹ yr⁻¹. This dimensionless prediction remains meaningful even when the metre is defined through c. A test requires actual frequency-ratio time series and a model of cavity aging, temperature, mirror response and instrumental drift. Published short-term clock precision alone is not an exclusion or a detection.')
p('A second priority is a full spacecraft range and Doppler fit with this specific affine history, moving endpoints, frequency ramps and clock standards. Earlier Cassini sensitivity comparisons and forecasts at reconstructed ranges do not constitute such a fit. Orbit and clock nuisance parameters can absorb parts of the proposed signal. Injection and recovery, followed by held-out arcs, are required before claiming compatibility or conflict.')
p('A third priority is joint distance, redshift and duration testing beyond the nearby sample, with calibration assumptions stated explicitly. For exact affine n and fixed source separation, redshift is constant with reception time, dz/dtᵣ=0. Canonical rolling χ instead gives dz/dtᵣ=γz at today’s normalization. Backreaction, source motion and gravity modify these predictions. Previously inspected reserved objects cannot be described as newly blind tests.')
p('The present result warrants inclusion as a candidate derivation with explicit limitations. It does not warrant presenting the coefficient, unchanged atomic behavior, global history, CMB origin or galaxy-rotation unification as established. Real-data tests are needed before stronger claims; discrepancies will be retained as branch-specific limitations rather than omitted.')
h('Additional references and reproducibility for Appendix D')
p('[18] Ortega-Gomez, A., et al. A tutorial on the conservation of momentum in photonic time-varying media (2023). https://arxiv.org/abs/2301.03333','Reference')
p('[19] Moussa, H., et al. Observation of Temporal Reflections and Broadband Frequency Translations at Photonic Time-Interfaces (2023). https://arxiv.org/abs/2208.07236','Reference')
p('[20] van de Bruck, C., Mifsud, J., and Nunes, N. J. The variation of the fine-structure constant from disformal couplings (2015). https://arxiv.org/abs/1510.00200','Reference')
p('Companion calculation: fixed_atom_time_field.zip, containing fixed_atom_time_field.md, calculate.py and results.json. The script uses Python, NumPy and SciPy. It checks ray stretching and conservative field–radiation dynamics; it contains no new observational likelihood.','Reference')
assert [p._p.xml for p in doc.paragraphs[:len(original)]]==original
doc.save(path)
print('Preserved',len(original),'original paragraphs; total',len(doc.paragraphs))
